from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from payroll.calculator import CoachBlock, PayrollLine, PayrollResult


def payroll_to_dict(result: PayrollResult) -> dict[str, Any]:
    def line_dict(line: PayrollLine) -> dict[str, Any]:
        return {"name": line.name, "qty": line.qty, "rate": line.rate, "amount": line.amount}

    return {
        "period": {"from": result.date_from, "to": result.date_to},
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "lesson_count": result.lesson_count,
        "grand_total": result.grand_total,
        "coaches": [
            {
                "coach": block.coach,
                "groups": [line_dict(l) for l in block.groups],
                "extras": [line_dict(l) for l in block.extras],
                "total": block.total,
            }
            for block in result.coaches
        ],
    }


def export_json(result: PayrollResult, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payroll_to_dict(result), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return path


def export_payroll_xlsx(result: PayrollResult, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Зарплата"
    ws.append(["Наименование", "Кол-во", "Ставка (руб)", "Сумма"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    row = 2
    total_cells: list[str] = []

    for block in result.coaches:
        ws.cell(row=row, column=1, value=block.coach).font = Font(bold=True)
        row += 1
        start = row

        if not block.groups:
            ws.cell(row=row, column=1, value="Нет занятий за период")
            row += 1
        else:
            for line in block.groups:
                ws.cell(row=row, column=1, value=line.name)
                ws.cell(row=row, column=2, value=line.qty)
                ws.cell(row=row, column=3, value=line.rate)
                ws.cell(row=row, column=4, value=f"=B{row}*C{row}")
                row += 1

        ws.cell(row=row, column=1, value="Доп.Расходы").font = Font(bold=True)
        row += 1
        for line in block.extras:
            ws.cell(row=row, column=1, value=line.name)
            ws.cell(row=row, column=2, value=line.qty)
            ws.cell(row=row, column=3, value=line.rate)
            ws.cell(row=row, column=4, value=f"=B{row}*C{row}")
            row += 1

        end = row - 1
        ws.cell(row=row, column=1, value="Итого").font = Font(bold=True)
        ws.cell(row=row, column=4, value=f"=SUM(D{start}:D{end})").font = Font(bold=True)
        total_cells.append(f"D{row}")
        row += 2

    ws.cell(row=row, column=1, value="ОБЩИЙ ИТОГ").font = Font(bold=True)
    if total_cells:
        ws.cell(row=row, column=4, value=f"=SUM({','.join(total_cells)})").font = Font(
            bold=True
        )

    ws.column_dimensions["A"].width = 56
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_tatyana_xlsx(
    rows: list[list],
    manual_rows: list[list],
    path: Path,
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Май"
    ws.append(["Тренер", "Статья ЗП", "Проект", "Сумма"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)

    ws2 = wb.create_sheet("для Татьяны (ручной ввод)")
    ws2.append(["Тренер", "Статья ЗП", "Примечание", "Сумма"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for row in manual_rows:
        ws2.append(row)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
